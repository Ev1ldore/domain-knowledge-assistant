"""Format-aware extraction. Page/section values are extracted, never invented."""
from pathlib import Path
import hashlib
import re
import zipfile

SUPPORTED={'.md','.txt','.pdf','.docx','.pptx','.csv'}


def extract(path):
    path=Path(path)
    ext=path.suffix.lower()
    if ext not in SUPPORTED: raise ValueError('不支持的文档格式')
    if ext in {'.docx','.pptx'}:
        with zipfile.ZipFile(path) as z:
            if sum(x.file_size for x in z.infolist())>100*1024*1024: raise ValueError('文档解压大小超限')
    if ext=='.pdf':
        import fitz
        with fitz.open(path) as doc:
            if len(doc)>1000: raise ValueError('PDF 页数超限')
            for i,page in enumerate(doc): yield page.get_text(),{'page':i+1}
    elif ext=='.docx':
        from docx import Document
        d=Document(path)
        for p in d.paragraphs:
            if p.text.strip(): yield p.text,{}
        for t in d.tables:
            yield '\n'.join(' | '.join(c.text for c in r.cells) for r in t.rows),{'section':'表格'}
    elif ext=='.pptx':
        from pptx import Presentation
        for i,slide in enumerate(Presentation(path).slides):
            yield '\n'.join(s.text for s in slide.shapes if s.has_text_frame),{'slide':i+1}
    elif ext=='.md':
        section=None; lines=[]
        for line in path.read_text(encoding='utf-8-sig').splitlines():
            if re.match(r'^#{1,6}\s+',line):
                if lines: yield '\n'.join(lines),({'section':section} if section else {})
                section=re.sub(r'^#+\s+','',line); lines=[line]
            else: lines.append(line)
        if lines: yield '\n'.join(lines),({'section':section} if section else {})
    else: yield path.read_text(encoding='utf-8-sig'),{}


def chunks(path,file_id,settings):
    from .evidence_store import stable_id
    version=hashlib.sha256(Path(path).read_bytes()).hexdigest()
    result=[]
    total=0
    for segment_index,(text,meta) in enumerate(extract(path)):
        text=text.strip()
        total+=len(text)
        if total>2_000_000: raise ValueError('解析文本超过 200 万字符限制')
        start=0
        while start<len(text):
            end=min(start+settings.chunk_size,len(text))
            if end<len(text):
                boundary=max(text.rfind('\n',start+settings.chunk_size//2,end),text.rfind('。',start+settings.chunk_size//2,end))
                if boundary>start: end=boundary+1
            fragment=text[start:end].strip()
            if fragment:
                cid=hashlib.sha256(f'{file_id}:{len(result)}:{fragment}'.encode()).hexdigest()[:24]
                offset=start+len(text[start:end])-len(text[start:end].lstrip())
                metadata={**meta,'document_version':version,'original_segment':text,
                    'locator':{**meta,'kind':'extracted_segment','segment_index':segment_index,
                    'start':offset,'end':offset+len(fragment),'segment_hash':stable_id(text)}}
                result.append({'id':cid,'text':fragment,'metadata':metadata})
            if end==len(text): break
            start=max(start+1,end-settings.chunk_overlap)
    if not result: raise ValueError('没有可提取文本；扫描件需要先 OCR，本版本不自动识别图片')
    return result
